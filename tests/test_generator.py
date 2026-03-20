"""Tests for the generator module (no LLM / network calls required)."""

import os
import tempfile

import pytest
from pptx import Presentation

from ppt_agente.generator import create_presentation


SAMPLE_SLIDES = [
    {
        "title": "Introduction to Python",
        "bullets": [
            "Python is a high-level language",
            "Easy to read and write",
            "Widely used in data science and web development",
        ],
    },
    {
        "title": "Key Features",
        "bullets": [
            "Dynamic typing",
            "Extensive standard library",
            "Cross-platform",
        ],
    },
    {
        "title": "Summary",
        "bullets": [
            "Python is powerful and beginner-friendly",
            "Great ecosystem of libraries",
            "Strong community support",
        ],
    },
]


def test_create_presentation_produces_file():
    with tempfile.TemporaryDirectory() as tmpdir:
        output = os.path.join(tmpdir, "test.pptx")
        result = create_presentation(SAMPLE_SLIDES, output_path=output)
        assert result == output
        assert os.path.isfile(output)


def test_create_presentation_slide_count():
    with tempfile.TemporaryDirectory() as tmpdir:
        output = os.path.join(tmpdir, "test.pptx")
        create_presentation(SAMPLE_SLIDES, output_path=output)
        prs = Presentation(output)
        assert len(prs.slides) == len(SAMPLE_SLIDES)


def test_create_presentation_titles():
    with tempfile.TemporaryDirectory() as tmpdir:
        output = os.path.join(tmpdir, "test.pptx")
        create_presentation(SAMPLE_SLIDES, output_path=output)
        prs = Presentation(output)
        for i, slide in enumerate(prs.slides):
            title_shape = slide.shapes.title
            assert title_shape is not None
            assert title_shape.text == SAMPLE_SLIDES[i]["title"]


def test_create_presentation_default_filename():
    original_dir = os.getcwd()
    with tempfile.TemporaryDirectory() as tmpdir:
        os.chdir(tmpdir)
        try:
            result = create_presentation(SAMPLE_SLIDES)
            assert result == "presentation.pptx"
            assert os.path.isfile("presentation.pptx")
        finally:
            os.chdir(original_dir)


def test_create_presentation_empty_slides():
    with tempfile.TemporaryDirectory() as tmpdir:
        output = os.path.join(tmpdir, "empty.pptx")
        create_presentation([], output_path=output)
        prs = Presentation(output)
        assert len(prs.slides) == 0
