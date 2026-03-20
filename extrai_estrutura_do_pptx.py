"""
extrai_estrutura_do_pptx.py — Funções utilitárias para extração e resumo de PPTX.
"""

from io import BytesIO

import streamlit as st
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from constantes_da_aplicacao import ACOES_LEITURA
from monitora_execucao_dos_agentes import monitor_agent


@monitor_agent("Extrator PPTX")
def extrair_estrutura_pptx(pptx_bytes: bytes) -> dict:
    """Extrai estrutura completa de um PPTX."""
    prs = Presentation(BytesIO(pptx_bytes))
    estrutura = {"total_slides": len(prs.slides), "slides": []}
    for idx_s, slide in enumerate(prs.slides):
        si = {"indice": idx_s, "numero": idx_s + 1,
              "layout": slide.slide_layout.name if slide.slide_layout else "?", "shapes": []}
        try:
            ns = slide.notes_slide
            si["notas"] = ns.notes_text_frame.text[:200] if ns.notes_text_frame.text else ""
        except Exception:
            si["notas"] = ""
        for idx_sh, shape in enumerate(slide.shapes):
            sh = {"indice": idx_sh, "nome": shape.name, "tipo": str(shape.shape_type)}
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                sh["eh_imagem"] = True
            if shape.has_chart:
                sh["eh_grafico"] = True
                sh["tipo_grafico"] = str(shape.chart.chart_type) if shape.chart else "?"
            if shape.has_text_frame:
                sh["text_frame"] = [
                    {"texto": p.text, "runs": [{"texto": r.text} for r in p.runs]}
                    for p in shape.text_frame.paragraphs
                ]
            if shape.has_table:
                tab = shape.table
                sh["tabela"] = {
                    "linhas": len(tab.rows), "colunas": len(tab.columns),
                    "celulas": [{"linha": ri, "coluna": ci, "texto": cell.text}
                                for ri, row in enumerate(tab.rows)
                                for ci, cell in enumerate(row.cells)]
                }
            si["shapes"].append(sh)
        estrutura["slides"].append(si)
    return estrutura


def gerar_resumo_estrutura(estrutura: dict) -> str:
    """Gera resumo textual da estrutura do PPTX para o LLM."""
    linhas = [f"Arquivo: {estrutura['total_slides']} slide(s)\n"]
    for sl in estrutura["slides"]:
        linhas.append(f"--- Slide {sl['numero']} (idx {sl['indice']}) | {sl['layout']} ---")
        for sh in sl["shapes"]:
            tipo_extra = ""
            if sh.get("eh_imagem"):
                tipo_extra = " [IMAGEM]"
            elif sh.get("eh_grafico"):
                tipo_extra = f" [GRÁFICO: {sh.get('tipo_grafico', '?')}]"
            linhas.append(f"  Shape {sh['indice']}: \"{sh['nome']}\"{tipo_extra}")
            if "text_frame" in sh:
                for pi, p in enumerate(sh["text_frame"]):
                    if p["texto"].strip():
                        linhas.append(f"    Par {pi}: \"{p['texto'][:200]}\"")
            if "tabela" in sh:
                t = sh["tabela"]
                linhas.append(f"    Tabela {t['linhas']}x{t['colunas']}")
                for c in t["celulas"]:
                    linhas.append(f"      [{c['linha']},{c['coluna']}]: \"{c['texto'][:120]}\"")
        if sl.get("notas"):
            linhas.append(f"  📝 Notas: \"{sl['notas'][:100]}\"")
        linhas.append("")
    return "\n".join(linhas)


def gerar_resumo_edit_log() -> str:
    """Gera resumo do log de alterações passadas para contexto do LLM."""
    if not st.session_state.edit_log:
        return ""
    linhas = ["ALTERAÇÕES JÁ REALIZADAS NESTA SESSÃO:"]
    for i, entry in enumerate(st.session_state.edit_log):
        linhas.append(
            f"  {i+1}. Slide {entry['slide']}, shape \"{entry['shape']}\": "
            f"'{entry['antes'][:50]}' → '{entry['depois'][:50]}' ({entry['acao']})"
        )
    return "\n".join(linhas)


def gerar_resumo_historico_alteracoes() -> str:
    """Gera resumo detalhado do historico_alteracoes para o LLM (v6)."""
    hist = st.session_state.historico_alteracoes
    if not hist:
        return ""
    linhas = ["HISTÓRICO DETALHADO DE ALTERAÇÕES NESTA SESSÃO:"]
    for i, h in enumerate(hist):
        linhas.append(
            f"  {i+1}. [{h.get('timestamp', '?')}] Instrução: \"{h.get('instrucao_original', '?')[:60]}\""
        )
        linhas.append(
            f"     Slide {h.get('slide', '?')}, shape idx {h.get('shape_idx', '?')}: "
            f"'{h.get('texto_antes', '')[:40]}' → '{h.get('texto_depois', '')[:40]}'"
        )
    return "\n".join(linhas)


def gerar_estrutura_formatada(pptx_bytes: bytes) -> str:
    """Gera exibição formatada da estrutura do PPTX para o chat."""
    prs = Presentation(BytesIO(pptx_bytes))
    linhas = [f"**📊 Estrutura do arquivo ({len(prs.slides)} slides):**\n"]
    for idx_s, slide in enumerate(prs.slides):
        titulo_slide = ""
        if slide.shapes.title and slide.shapes.title.has_text_frame:
            titulo_slide = slide.shapes.title.text_frame.text[:60]
        layout = slide.slide_layout.name if slide.slide_layout else "?"
        linhas.append(f"**📄 Slide {idx_s + 1}** — *{titulo_slide or layout}*")

        for idx_sh, shape in enumerate(slide.shapes):
            if shape.has_table:
                tab = shape.table
                tipo = f"Tabela {len(tab.rows)}x{len(tab.columns)}"
                texto = ""
            elif hasattr(shape, 'has_chart') and shape.has_chart:
                tipo = f"Gráfico ({shape.chart.chart_type})" if shape.chart else "Gráfico"
                texto = ""
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                tipo = "Imagem"
                texto = ""
            elif shape.has_text_frame:
                tipo = "Título" if shape == slide.shapes.title else "Caixa de texto"
                texto = shape.text_frame.text[:80].replace("\n", " ")
            else:
                tipo = str(shape.shape_type).replace("MSO_SHAPE_TYPE.", "")
                texto = ""

            if texto:
                texto_safe = texto.replace("<", "&lt;").replace(">", "&gt;")
                linhas.append(f"&nbsp;&nbsp;&nbsp;&nbsp;└── Shape {idx_sh} | {tipo} | `{texto_safe}`")
            else:
                linhas.append(f"&nbsp;&nbsp;&nbsp;&nbsp;└── Shape {idx_sh} | {tipo} | *(sem texto)*")

        try:
            ns = slide.notes_slide
            notas = ns.notes_text_frame.text
            if notas and notas.strip():
                notas_safe = notas[:60].replace("<", "&lt;").replace(">", "&gt;")
                linhas.append(f"&nbsp;&nbsp;&nbsp;&nbsp;📝 Notas: *{notas_safe}*")
        except Exception:
            pass

        linhas.append("")
    return "\n\n".join(linhas)


def _eh_acao_leitura(comando: dict) -> bool:
    """Verifica se um comando é de leitura (não altera o PPTX)."""
    if comando is None:
        return False
    acao = comando.get("acao", "")
    tipo = comando.get("tipo", "")
    if tipo == "leitura":
        return True
    if acao in ACOES_LEITURA:
        return True
    return False
